import json
import os
from fnmatch import fnmatch
from pathlib import Path


def _load_search_roots() -> list[str]:
    try:
        profile = json.loads(
            (Path(__file__).parent.parent / "profile.json").read_text(encoding="utf-8")
        )
        roots = [
            profile.get("desktop", ""),
            profile.get("documents", ""),
            profile.get("downloads", ""),
            profile.get("onedrive", ""),
        ]
        return [r for r in roots if r]
    except Exception:
        home = str(Path.home())
        return [
            os.path.join(home, "OneDrive", "Desktop"),
            os.path.join(home, "OneDrive", "Documents"),
            os.path.join(home, "Downloads"),
            os.path.join(home, "OneDrive"),
        ]


_SKIP_DIRS = {
    "node_modules", "__pycache__", ".git", ".svn",
    "AppData", "Windows", "Program Files", "Program Files (x86)",
    "$Recycle.Bin", "System Volume Information",
}


def search_files(pattern: str, folder: str = None) -> str:
    """
    Search recursively for files whose names contain `pattern`.
    Returns a newline-separated list of absolute paths (up to 20 results).
    """
    roots = [folder] if folder else _load_search_roots()
    matches: list[str] = []

    for root in roots:
        if not root or not os.path.exists(root):
            continue

        for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
            dirnames[:] = [
                d for d in dirnames
                if d not in _SKIP_DIRS and not d.startswith(".")
            ]

            for fname in filenames:
                if fnmatch(fname.lower(), f"*{pattern.lower()}*"):
                    matches.append(os.path.join(dirpath, fname))
                    if len(matches) >= 20:
                        return "\n".join(matches)

    if not matches:
        return f"No files found matching '{pattern}'"
    return "\n".join(matches)


def open_file(path: str) -> str:
    """Open a file with its default Windows application via os.startfile."""
    if not os.path.exists(path):
        return f"File not found: {path}"
    try:
        os.startfile(path)
        return f"Opened: {os.path.basename(path)}"
    except Exception as e:
        return f"Failed to open file: {e}"


# ── Document readers ───────────────────────────────────────────────────────────

def _read_pdf(path: str, max_chars: int) -> str:
    try:
        import fitz  # pymupdf
    except ImportError:
        return "PDF reading unavailable — run: pip install pymupdf"
    try:
        doc = fitz.open(path)
        total_pages = len(doc)
        parts = []
        chars = 0
        last_page = 0
        for i, page in enumerate(doc):
            text = page.get_text()
            if not text.strip():
                continue
            chunk = f"[Page {i + 1}]\n{text}\n"
            if chars + len(chunk) > max_chars:
                # Include partial page up to limit
                remaining = max_chars - chars
                parts.append(chunk[:remaining])
                last_page = i + 1
                break
            parts.append(chunk)
            chars += len(chunk)
            last_page = i + 1
        doc.close()
        header = f"[PDF: {os.path.basename(path)} — {total_pages} pages total, showing pages 1-{last_page}]\n\n"
        body = "".join(parts)
        suffix = "\n[... truncated — ask to continue from a specific page ...]" if last_page < total_pages else ""
        return header + body + suffix
    except Exception as e:
        return f"Cannot read PDF: {e}"


def _read_docx(path: str, max_chars: int) -> str:
    try:
        from docx import Document
    except ImportError:
        return "Word reading unavailable — run: pip install python-docx"
    try:
        doc = Document(path)
        parts = []
        chars = 0
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ""
            prefix = "## " if style.startswith("Heading") else ""
            line = f"{prefix}{text}\n"
            if chars + len(line) > max_chars:
                parts.append(line[:max_chars - chars])
                break
            parts.append(line)
            chars += len(line)
        header = f"[Word document: {os.path.basename(path)}]\n\n"
        body = "".join(parts)
        suffix = "\n[... truncated ...]" if chars >= max_chars else ""
        return header + body + suffix
    except Exception as e:
        return f"Cannot read Word file: {e}"


def _read_xlsx(path: str, max_chars: int) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "Excel reading unavailable — run: pip install openpyxl"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        chars = 0
        truncated = False
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            section = f"Sheet: {sheet_name}\n"
            if chars + len(section) > max_chars:
                truncated = True
                break
            parts.append(section)
            chars += len(section)
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = " | ".join("" if v is None else str(v) for v in row)
                if not cells.strip():
                    continue
                line = f"  Row {i}: {cells}\n"
                if chars + len(line) > max_chars:
                    truncated = True
                    break
                parts.append(line)
                chars += len(line)
            if truncated:
                break
        wb.close()
        header = f"[Excel file: {os.path.basename(path)} — sheets: {', '.join(wb.sheetnames)}]\n\n"
        body = "".join(parts)
        suffix = "\n[... truncated ...]" if truncated else ""
        return header + body + suffix
    except Exception as e:
        return f"Cannot read Excel file: {e}"


# ── Extension dispatch map ─────────────────────────────────────────────────────

_EXT_READERS = {
    ".pdf":  _read_pdf,
    ".docx": _read_docx,
    ".doc":  _read_docx,
    ".xlsx": _read_xlsx,
    ".xls":  _read_xlsx,
}


def read_file_content(path: str, max_chars: int = 12000) -> str:
    """Read content from a file. Supports PDF, Word (.docx), Excel (.xlsx), and plain text."""
    if not os.path.exists(path):
        return f"File not found: {path}"

    ext = Path(path).suffix.lower()
    reader = _EXT_READERS.get(ext)
    if reader:
        return reader(path, max_chars)

    # Plain text fallback (UTF-8 with replacement for bad bytes)
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read(max_chars)
        if len(content) == max_chars:
            content += "\n[... truncated — file has more content ...]"
        return content
    except Exception as e:
        return f"Cannot read file: {e}"
